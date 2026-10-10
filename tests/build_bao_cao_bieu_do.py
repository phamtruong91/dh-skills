"""Thử báo cáo nghiên cứu / số liệu có biểu đồ (dữ liệu GIẢ, cố định, lặp lại được).

3 báo cáo: phan-tich-ket-qua-khao-sat (Word + Excel có biểu đồ gốc), bao-cao-tien-do-de-tai, bao-cao-khcn-nam.
Đóng vai "AI làm theo SKILL.md" (không phải AI độc lập). Biểu đồ vẽ bằng matplotlib, chèn PNG vào Word;
văn bản alt của hình ghi lại đúng số liệu để bộ kiểm tra đối chiếu với bảng.
Dùng: python -X utf8 tests/build_bao_cao_bieu_do.py
"""
import json
import sys
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
from matplotlib import font_manager  # noqa: E402

sys.path.insert(0, str(Path(__file__).resolve().parent))
from build_tai_lon import *  # noqa: F401,F403,E402
from build_tai_lon import C, J, L, NOTE, CQ, TR, dump, grid, head, keep_together, money, page_number, sign  # noqa: E402
from build_reference_outputs import DOT, O, S, new_doc, para  # noqa: E402

HERE = Path(__file__).resolve().parent
R = HERE / "results"
IMG = HERE / "outputs" / "hinh"
IMG.mkdir(parents=True, exist_ok=True)
for _f in ("LiberationSerif-Regular.ttf", "LiberationSerif-Bold.ttf", "LiberationSerif-Italic.ttf"):
    _p = Path("/usr/share/fonts/truetype/liberation") / _f
    if _p.exists():
        font_manager.fontManager.addfont(str(_p))
plt.rcParams.update({"font.family": "Liberation Serif", "font.size": 10, "axes.spines.top": False, "axes.spines.right": False,
                     "axes.axisbelow": True, "axes.grid": True, "grid.color": "#d9d9d9", "grid.linewidth": 0.6})
MAU = ["#1F3864", "#8FAADC", "#D9E2F3"]  # đậm / vừa / nhạt: vẫn phân biệt được khi in đen trắng
HATCH = ["", "///", "xxx"]
LIKERT = ["#EAF0FA", "#BFD0EE", "#8FAADC", "#4F76B8", "#1F3864"]


def grid_k(d, headers, rows, widths, **kw):
    """Bảng ngắn (đến 12 dòng) giữ nguyên trên một trang: mọi đoạn trong bảng dính với dòng sau."""
    t = grid(d, headers, rows, widths, **kw)
    if len(rows) <= 12:
        for r in t.rows[:-1]:
            for c in r.cells:
                for p_ in c.paragraphs:
                    p_.paragraph_format.keep_with_next = True
    return t


def vn(x, nd=2):
    """Số thập phân kiểu Việt Nam (dấu phẩy)."""
    return f"{x:.{nd}f}".replace(".", ",")


def pv(s):
    return float(str(s).replace(".", "").replace(",", ".")) if "," in str(s) and str(s).count(",") == 1 and "." not in str(s).split(",")[0][-3:-2] else float(str(s).replace(",", "."))


def save(fig, name):
    p = IMG / f"{name}.png"
    fig.savefig(p, dpi=220, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    return p


def add_figure(d, png, no, title, source, alt, note=None):
    """Hình + chú thích 'Hình n. ...' dưới hình + dòng nguồn; ba đoạn dính nhau để không tách trang."""
    p = d.add_paragraph()
    p.alignment = C
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run()
    run.add_picture(str(png), width=Cm(15.0))
    for dp in run._r.xpath(".//wp:docPr"):
        dp.set("descr", alt)
        dp.set("title", f"Hình {no}")
    para(d, f"Hình {no}. {title}", bold=True, size=13, align=C, after=0, keep=True)
    para(d, f"Nguồn: {source}" + (f". {note}" if note else ""), italic=True, size=13, align=C, after=8)


def alt_of(kind, pairs, unit=""):
    return f"{kind}. Số liệu: " + "; ".join(f"{k}={v}" for k, v in pairs) + (f". Đơn vị: {unit}" if unit else "")


# =====================================================================
# 1. phan-tich-ket-qua-khao-sat
# =====================================================================
CAU = [  # (nhóm, nội dung, đảo?, phân bố mức 1..5)
    ("Giảng dạy", "Nội dung bài giảng rõ ràng, dễ hiểu", False, [10, 28, 110, 360, 334]),
    ("Giảng dạy", "Giảng viên nhiệt tình, tận tâm", False, [6, 20, 90, 330, 396]),
    ("Giảng dạy", "Kiểm tra, đánh giá công bằng", False, [22, 60, 200, 330, 230]),
    ("Giảng dạy", "Tài liệu học tập đầy đủ", False, [18, 70, 230, 310, 214]),
    ("Cơ sở vật chất", "Phòng học đủ tiện nghi", False, [60, 140, 260, 250, 132]),
    ("Cơ sở vật chất", "Mạng wifi ổn định", False, [150, 210, 240, 170, 72]),
    ("Cơ sở vật chất", "Thư viện đáp ứng nhu cầu", False, [40, 90, 260, 280, 160]),  # tổng 830: 12 phiếu bỏ trống câu này
    ("Cơ sở vật chất", "Phòng thực hành đủ thiết bị", False, [80, 150, 250, 230, 132]),
    ("Hỗ trợ sinh viên", "Thủ tục hành chính rườm rà (câu đảo)", True, [70, 180, 220, 230, 142]),
    ("Hỗ trợ sinh viên", "Tư vấn học tập kịp thời", False, [24, 66, 220, 320, 212]),
    ("Hỗ trợ sinh viên", "Hỗ trợ tìm việc làm hiệu quả", False, [96, 170, 300, 190, 86]),
    ("Hỗ trợ sinh viên", "Thắc mắc được giải đáp nhanh", False, [14, 40, 170, 350, 268]),
]
TB_KY_TRUOC = [4.02, 4.18, 3.70, 3.55, 3.10, 2.60, 3.40, 3.05, 3.40, 3.62, 3.01, 3.90]  # Q5 đổi nội dung câu hỏi; Q9 là điểm gốc chưa đảo
KHOA_KS = [("Công nghệ thông tin", 300, 3.71), ("Kinh tế", 280, 3.52), ("Ngoại ngữ", 190, 3.38), ("Sư phạm", 48, 3.60), ("Luật", 24, 4.20)]


def mean(dist, rev=False):
    n = sum(dist)
    s = sum((6 - (k + 1) if rev else k + 1) * v for k, v in enumerate(dist))
    return s / n


def xep_loai(m):
    return "Tốt" if m >= 4.0 else ("Trung bình" if m >= 3.0 else "Cần cải thiện")


def gen_khao_sat():
    cau = [{"stt": k + 1, "nhom": g, "noi_dung": t, "dao": dao, "phan_bo_muc_1_den_5": dist,
            "diem_tb_khai_bao": (4.05 if k == 2 else round(mean(dist), 2)), "tb_ky_truoc": TB_KY_TRUOC[k],
            "cau_hoi_doi_noi_dung_so_voi_ky_truoc": k == 4} for k, (g, t, dao, dist) in enumerate(CAU)]
    return {"co_quan_chu_quan": CQ, "co_quan_ban_hanh": TR + " – PHÒNG KHẢO THÍ VÀ ĐẢM BẢO CHẤT LƯỢNG", "dia_danh": "Hà Nội",
            "so_van_ban": None, "doi_tuong": "Sinh viên", "muc_dich_khao_sat": "Đánh giá mức độ hài lòng của sinh viên về hoạt động đào tạo và hỗ trợ, học kỳ 1 năm học 2026–2027",
            "so_phieu_phat_ra": 1200, "so_phieu_thu_ve": 1010, "so_phieu_hop_le": 842,
            "ly_do_loai": {"bo_trong_tren_30_phan_tram": 91, "chon_mot_dap_an_cho_ca_bang": 52, "mau_thuan_logic": 27},
            "thang_do": "Likert 5 mức (1 = rất không hài lòng ... 5 = rất hài lòng)", "du_lieu_tong_hop": cau,
            "theo_khoa": [{"khoa": k, "so_phieu": n, "diem_tb_chung_khai_bao": m} for k, n, m in KHOA_KS],
            "y_kien_mo": {"so_phieu_co_tra_loi_mo": 360, "chu_de": [["Mạng wifi yếu, chập chờn", 120], ["Thủ tục hành chính chậm", 85],
                                                                       ["Giảng viên nhiệt tình", 95], ["Thiếu chỗ ngồi giờ cao điểm", 40]],
                          "trich_dan": ["Wifi ở giảng đường B hay mất kết nối vào giờ cao điểm.", "Thầy cô hướng dẫn rất tận tình."]},
            "ky_truoc": "học kỳ 1 năm học 2025–2026", "nguoi_ky": {"chuc_danh": "Trưởng phòng Khảo thí và Đảm bảo chất lượng", "ho_ten": None},
            "noi_nhan": ["Ban Giám hiệu", "Các khoa, các đơn vị liên quan", "Lưu: VT, KT&ĐBCL"]}


def build_khao_sat():
    inp = gen_khao_sat()
    cau = inp["du_lieu_tong_hop"]
    n_valid = inp["so_phieu_hop_le"]
    rows, tb_list, n_ans = [], [], []
    for c in cau:
        dist = c["phan_bo_muc_1_den_5"]
        m = mean(dist, c["dao"])
        tb_list.append(m)
        n_ans.append(sum(dist))
    tl_phan_hoi = n_valid / inp["so_phieu_phat_ra"] * 100
    # nhóm
    nhom_names = ["Giảng dạy", "Cơ sở vật chất", "Hỗ trợ sinh viên"]
    nhom_tb = {g: sum(tb_list[k] for k, c in enumerate(cau) if c["nhom"] == g) / sum(1 for c in cau if c["nhom"] == g) for g in nhom_names}
    tb_chung = sum(tb_list) / len(tb_list)
    # so với kỳ trước (chỉ câu so sánh được): Q5 đổi nội dung -> bỏ; Q9 đối chiếu điểm đã đảo
    so_sanh = []
    for k, c in enumerate(cau):
        if c["cau_hoi_doi_noi_dung_so_voi_ky_truoc"]:
            so_sanh.append(None)
            continue
        prev = 6 - c["tb_ky_truoc"] if c["dao"] else c["tb_ky_truoc"]
        so_sanh.append(prev)
    lech_khai_bao = [(k + 1, c["diem_tb_khai_bao"], tb_list[k]) for k, c in enumerate(cau) if abs(c["diem_tb_khai_bao"] - mean(c["phan_bo_muc_1_den_5"])) > 0.005 and not c["dao"]]
    q9_raw = mean(cau[8]["phan_bo_muc_1_den_5"])
    khoa_ok = [x for x in KHOA_KS if x[1] >= 30]
    khoa_nho = [x for x in KHOA_KS if x[1] < 30]
    tb_khoa_gq = sum(n * m for _, n, m in KHOA_KS) / sum(n for _, n, _ in KHOA_KS)
    tong_loai = sum(inp["ly_do_loai"].values())
    loai_thuc = inp["so_phieu_thu_ve"] - n_valid

    # ---- biểu đồ ----
    # Hình 1: điểm TB từng câu, ngưỡng 3,0 và 4,0
    fig, ax = plt.subplots(figsize=(6.4, 3.7))
    labels = [f"Câu {k + 1}" for k in range(len(cau))]
    cols = ["#1F3864" if m >= 4 else ("#8FAADC" if m >= 3 else "#D9E2F3") for m in tb_list]
    hat = ["" if m >= 4 else ("///" if m >= 3 else "xxx") for m in tb_list]
    bars = ax.bar(labels, tb_list, color=cols, edgecolor="#1F3864", linewidth=0.8)
    for b, h in zip(bars, hat):
        b.set_hatch(h)
    for b, m in zip(bars, tb_list):
        ax.text(b.get_x() + b.get_width() / 2, m + 0.05, vn(m), ha="center", va="bottom", fontsize=8.5)
    ax.axhline(4.0, color="#444", ls="--", lw=0.9, label="Ngưỡng 4,0 (từ mức này là tốt)")
    ax.axhline(3.0, color="#444", ls=":", lw=1.1, label="Ngưỡng 3,0 (dưới mức này là cần cải thiện)")
    ax.legend(loc="upper center", bbox_to_anchor=(0.5, -0.2), ncol=1, frameon=False, fontsize=8.5)
    ax.set_ylim(0, 5)
    ax.set_ylabel("Điểm trung bình (thang 1–5)")
    ax.tick_params(axis="x", labelrotation=45, labelsize=8.5)
    ax.yaxis.set_major_formatter(matplotlib.ticker.FuncFormatter(lambda v, _: vn(v, 1)))
    p1 = save(fig, "khao-sat-hinh1")
    # Hình 2: phân bố % theo mức (100%)
    fig, ax = plt.subplots(figsize=(6.4, 3.9))
    ys = list(range(len(cau)))[::-1]
    left = [0.0] * len(cau)
    for lv in range(5):
        vals = [100 * c["phan_bo_muc_1_den_5"][lv] / n_ans[k] for k, c in enumerate(cau)]
        ax.barh(ys, vals, left=left, color=LIKERT[lv], edgecolor="white", linewidth=0.6, label=f"Mức {lv + 1}")
        for y, v, l_ in zip(ys, vals, left):
            if v >= 9:
                ax.text(l_ + v / 2, y, f"{v:.0f}".replace(".", ","), ha="center", va="center", fontsize=7.5, color="white" if lv >= 3 else "black")
        left = [a + b for a, b in zip(left, vals)]
    ax.set_yticks(ys)
    ax.set_yticklabels([f"Câu {k + 1}" for k in range(len(cau))], fontsize=8.5)
    ax.set_xlim(0, 100)
    ax.set_xlabel("Tỷ lệ phiếu trả lời (%)")
    ax.legend(ncol=5, loc="upper center", bbox_to_anchor=(0.5, -0.16), frameon=False, fontsize=8.5)
    ax.grid(axis="y", visible=False)
    p2 = save(fig, "khao-sat-hinh2")
    # Hình 3: kỳ này so với kỳ trước (chỉ câu so sánh được)
    idx = [k for k, v in enumerate(so_sanh) if v is not None]
    fig, ax = plt.subplots(figsize=(6.4, 3.4))
    w = 0.38
    xs = list(range(len(idx)))
    b1 = ax.bar([x - w / 2 for x in xs], [so_sanh[k] for k in idx], w, color=MAU[1], hatch="///", edgecolor="#1F3864", label="Học kỳ 1 năm học 2025–2026")
    b2 = ax.bar([x + w / 2 for x in xs], [tb_list[k] for k in idx], w, color=MAU[0], edgecolor="#1F3864", label="Học kỳ 1 năm học 2026–2027")
    for bs in (b1, b2):
        for b in bs:
            ax.text(b.get_x() + b.get_width() / 2, b.get_height() + 0.04, vn(b.get_height()), ha="center", fontsize=6, rotation=90)
    ax.set_xticks(xs)
    ax.set_xticklabels([f"Câu {k + 1}" for k in idx], fontsize=8.5, rotation=45)
    ax.set_ylim(0, 5.6)
    ax.set_ylabel("Điểm trung bình (thang 1–5)")
    ax.yaxis.set_major_formatter(matplotlib.ticker.FuncFormatter(lambda v, _: vn(v, 1)))
    ax.legend(loc="upper center", bbox_to_anchor=(0.5, -0.24), ncol=2, frameon=False, fontsize=8.5)
    p3 = save(fig, "khao-sat-hinh3")
    # Hình 4: điểm TB chung theo khoa (không đưa nhóm dưới 30 phiếu)
    fig, ax = plt.subplots(figsize=(6.4, 2.9))
    nm = [f"{k}\n(n = {n})" for k, n, _ in khoa_ok]
    vv = [m for _, _, m in khoa_ok]
    bs = ax.bar(nm, vv, color=MAU[1], edgecolor="#1F3864")
    for b, v in zip(bs, vv):
        ax.text(b.get_x() + b.get_width() / 2, v + 0.05, vn(v), ha="center", fontsize=9)
    ax.set_ylim(0, 5)
    ax.set_ylabel("Điểm trung bình chung (thang 1–5)")
    ax.yaxis.set_major_formatter(matplotlib.ticker.FuncFormatter(lambda v, _: vn(v, 1)))
    ax.tick_params(axis="x", labelsize=8.5)
    p4 = save(fig, "khao-sat-hinh4")

    # ---- Word ----
    d = new_doc()
    page_number(d)
    head(d, inp, "BC")
    para(d, "BÁO CÁO", bold=True, size=14, align=C, after=0)
    para(d, "PHÂN TÍCH KẾT QUẢ KHẢO SÁT MỨC ĐỘ HÀI LÒNG CỦA SINH VIÊN", bold=True, align=C, after=0)
    para(d, "Học kỳ 1 năm học 2026–2027", bold=True, align=C, after=8)
    para(d, "I. THÔNG TIN CHUNG", bold=True, after=4, keep=True)
    para(d, f"Mục đích: {inp['muc_dich_khao_sat']}. Thang đo: {inp['thang_do']}.", first_indent=1.0, align=J)
    para(d, f"Số phiếu phát ra {money(inp['so_phieu_phat_ra'])}, thu về {money(inp['so_phieu_thu_ve'])}, hợp lệ {n_valid}; tỷ lệ phiếu hợp lệ trên phiếu phát ra là "
            f"{vn(tl_phan_hoi, 1)}%, đạt yêu cầu tối thiểu 60%.", first_indent=1.0, align=J)
    para(d, "Bảng 1. Cơ cấu phiếu khảo sát", bold=True, size=13, align=C, after=2, keep=True)
    ly = inp["ly_do_loai"]
    grid_k(d, ["STT", "Chỉ tiêu", "Số phiếu"], [["1", "Phiếu phát ra", money(inp["so_phieu_phat_ra"])], ["2", "Phiếu thu về", money(inp["so_phieu_thu_ve"])],
                                                  ["3", "Loại: bỏ trống trên 30% số câu", str(ly["bo_trong_tren_30_phan_tram"])],
                                                  ["4", "Loại: chọn một đáp án cho cả bảng", str(ly["chon_mot_dap_an_cho_ca_bang"])],
                                                  ["5", "Loại: mâu thuẫn logic", str(ly["mau_thuan_logic"])], ["6", "Phiếu hợp lệ", str(n_valid)]],
         [1.6, 10.4, 4.0], aligns=[C, L, C])
    para(d, f"Lưu ý: tổng số phiếu loại theo ba lý do trong dữ liệu đầu vào là {tong_loai}, trong khi phiếu thu về trừ phiếu hợp lệ là {loai_thuc}; hai con số lệch "
            f"{tong_loai - loai_thuc} phiếu. Báo cáo giữ nguyên số liệu đầu vào, chưa tự điều chỉnh; đề nghị Phòng rà lại danh sách phiếu loại trước khi ký.",
         italic=True, size=13, first_indent=1.0, align=J, after=8)
    para(d, "II. KẾT QUẢ CHI TIẾT", bold=True, after=4, keep=True)
    para(d, "Điểm trung bình tính trên số phiếu có trả lời từng câu. Câu 9 là câu đảo nên đã mã hóa ngược (điểm = 6 − mức chọn) trước khi tính. "
            "Quy ước xếp loại: từ 4,0 trở lên là tốt; từ 3,0 đến 3,99 là trung bình; dưới 3,0 là cần cải thiện.", first_indent=1.0, align=J)
    add_figure(d, p1, 1, "Điểm trung bình từng câu hỏi và ngưỡng xếp loại", "Dữ liệu khảo sát học kỳ 1 năm học 2026–2027, Phòng Khảo thí và Đảm bảo chất lượng",
               alt_of("Biểu đồ cột điểm trung bình theo câu", [(f"Câu {k + 1}", vn(m)) for k, m in enumerate(tb_list)], "điểm, thang 1–5"),
               note="Câu 9 đã mã hóa ngược")
    rows2, cap = [], []
    for k, c in enumerate(cau):
        prev = "" if so_sanh[k] is None else vn(so_sanh[k])
        diff = "" if so_sanh[k] is None else ("+" if tb_list[k] - so_sanh[k] >= 0 else "−") + vn(abs(tb_list[k] - so_sanh[k]))
        rows2.append([str(k + 1), c["noi_dung"], str(n_ans[k]), vn(tb_list[k]), xep_loai(tb_list[k]), prev, diff])
    para(d, "Bảng 2. Điểm trung bình từng câu hỏi và so sánh với kỳ trước", bold=True, size=13, align=C, after=2, keep=True)
    grid_k(d, ["STT", "Nội dung câu hỏi", "Số trả lời", "Điểm TB", "Xếp loại", "Kỳ trước", "Chênh lệch"], rows2, [1.4, 5.4, 1.7, 1.6, 2.3, 1.7, 1.9],
         total=["", "Điểm trung bình chung (12 câu)", "", vn(tb_chung), xep_loai(tb_chung), "", ""], aligns=[C, L, C, C, C, C, C], size=11)
    para(d, "Ghi chú: câu 5 đổi nội dung so với kỳ trước nên không so sánh (để trống). Điểm kỳ trước của câu 9 đã mã hóa ngược để cùng chiều với kỳ này.",
         italic=True, size=13, first_indent=1.0, align=J, after=8)
    add_figure(d, p2, 2, "Phân bố mức trả lời từng câu hỏi (%)", "Dữ liệu khảo sát học kỳ 1 năm học 2026–2027",
               alt_of("Biểu đồ cột xếp chồng 100% phân bố mức 1–5", [(f"Câu {k + 1}", "/".join(f"{100 * v / n_ans[k]:.1f}" for v in c['phan_bo_muc_1_den_5'])) for k, c in enumerate(cau)], "% phiếu trả lời"),
               note="Câu 9 giữ nguyên mức người trả lời chọn, chưa mã hóa ngược; tỷ lệ tính trên số phiếu có trả lời câu đó")
    para(d, "III. NHẬN XÉT", bold=True, after=4, keep=True)
    order = sorted(range(len(cau)), key=lambda k: -tb_list[k])
    top = order[:3]
    bot = order[-3:][::-1]
    para(d, "1. Điểm mạnh: " + "; ".join(f"câu {k + 1} ({cau[k]['noi_dung'].lower()}) đạt {vn(tb_list[k])}" for k in top) + ".", first_indent=1.0, align=J)
    para(d, "2. Điểm cần cải thiện nhất: " + "; ".join(f"câu {k + 1} ({cau[k]['noi_dung'].lower()}) chỉ đạt {vn(tb_list[k])}" for k in bot) + ".", first_indent=1.0, align=J)
    para(d, "3. Theo nhóm nội dung: " + "; ".join(f"{g} {vn(nhom_tb[g])}" for g in nhom_names) + f". Điểm trung bình chung {vn(tb_chung)}, xếp loại {xep_loai(tb_chung).lower()}.",
         first_indent=1.0, align=J)
    ht = [(k, tb_list[k] - so_sanh[k]) for k in idx]
    ht.sort(key=lambda x: x[1])
    para(d, f"4. So với kỳ trước: cải thiện nhiều nhất là câu {ht[-1][0] + 1} (+{vn(ht[-1][1])}); giảm nhiều nhất là câu {ht[0][0] + 1} ({vn(ht[0][1]).replace('-', '−')}). Chỉ so sánh các câu giữ nguyên nội dung.",
         first_indent=1.0, align=J, keep=True)
    add_figure(d, p3, 3, "So sánh điểm trung bình giữa hai kỳ khảo sát (các câu giữ nguyên nội dung)", "Dữ liệu khảo sát hai học kỳ 1 năm học 2025–2026 và 2026–2027",
               alt_of("Biểu đồ cột ghép kỳ trước và kỳ này", [(f"Câu {k + 1}", f"{vn(so_sanh[k])}->{vn(tb_list[k])}") for k in idx], "điểm, thang 1–5"))
    para(d, "5. Theo khoa: " + "; ".join(f"{k} {vn(m)} ({n} phiếu)" for k, n, m in khoa_ok) + ". "
         + "; ".join(f"Khoa {k} chỉ có {n} phiếu (dưới 30) nên không kết luận riêng" for k, n, _ in khoa_nho) + ".", first_indent=1.0, align=J, keep=True)
    add_figure(d, p4, 4, "Điểm trung bình chung theo khoa (các khoa từ 30 phiếu trở lên)", "Dữ liệu khảo sát học kỳ 1 năm học 2026–2027, điểm trung bình chung do đơn vị cung cấp",
               alt_of("Biểu đồ cột điểm trung bình chung theo khoa", [(k, vn(m)) for k, n, m in khoa_ok], "điểm, thang 1–5"))
    y = inp["y_kien_mo"]
    para(d, f"6. Ý kiến mở ({y['so_phieu_co_tra_loi_mo']} phiếu có trả lời): " + "; ".join(f"{t} ({c} ý kiến, {vn(100 * c / y['so_phieu_co_tra_loi_mo'], 1)}%)" for t, c in y["chu_de"]) + ". "
         f"Trích dẫn tiêu biểu (đã ẩn danh): “{y['trich_dan'][0]}”; “{y['trich_dan'][1]}”.", first_indent=1.0, align=J)
    para(d, "IV. HẠN CHẾ CỦA SỐ LIỆU", bold=True, after=4, keep=True)
    lim = [f"Số phiếu loại theo lý do và theo đối chiếu thu về/hợp lệ lệch {tong_loai - loai_thuc} phiếu (mục I).",
           f"Câu 7 có {n_valid - n_ans[6]} phiếu bỏ trống nên điểm trung bình tính trên {n_ans[6]} phiếu; không điền hộ câu trả lời.",
           f"Câu 3: điểm trung bình do đơn vị cung cấp là {vn(cau[2]['diem_tb_khai_bao'])}, tính lại từ phân bố mức là {vn(tb_list[2])}; báo cáo dùng giá trị tính lại.",
           f"Điểm trung bình chung theo khoa (gia quyền) là {vn(tb_khoa_gq)}, khác điểm trung bình chung tính từ 12 câu ({vn(tb_chung)}) {vn(abs(tb_khoa_gq - tb_chung))} điểm; hai nguồn tính khác nhau, cần đối chiếu trước khi công bố điểm theo khoa.",
           "Tỷ lệ hợp lệ trên phiếu phát ra đạt yêu cầu 60%, nhưng chưa có thông tin về mức đại diện của mẫu theo khoa, khóa."]
    for k, t in enumerate(lim, 1):
        para(d, f"{k}. {t}", first_indent=1.0, align=J)
    para(d, "V. ĐỀ XUẤT CẢI TIẾN", bold=True, after=4, keep=True)
    de_xuat = [(bot[0], "Rà soát chất lượng mạng wifi tại các giảng đường", "Đơn vị phụ trách mạng và thiết bị"),
               (bot[1], "Rà soát và rút gọn thủ tục hành chính sinh viên thường dùng", "Phòng Đào tạo, Phòng Công tác sinh viên"),
               (bot[2], "Tăng hoạt động kết nối việc làm cho sinh viên", "Đơn vị phụ trách quan hệ doanh nghiệp")]
    for k, (q, t, dv) in enumerate(de_xuat, 1):
        para(d, f"{k}. {t} (căn cứ câu {q + 1}, điểm {vn(tb_list[q])}). Đơn vị đề xuất phối hợp: {dv} (chờ lãnh đạo xác nhận). Thời hạn: {DOT}.", first_indent=1.0, align=J)
    keep_together(d, 2)
    para(d, "", after=6)
    sign(d, inp["noi_nhan"], inp["nguoi_ky"]["chuc_danh"], inp["nguoi_ky"]["ho_ten"])
    out = O / "phan-tich-ket-qua-khao-sat.docx"
    d.save(out)

    # ---- Excel: bảng số liệu + biểu đồ gốc (sửa được) ----
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, Reference
    from openpyxl.styles import Font, PatternFill
    from openpyxl.worksheet.properties import PageSetupProperties
    wb = Workbook()
    ws = wb.active
    ws.title = "Du_lieu"
    ws.append(["STT", "Nội dung", "Nhóm", "Câu đảo", "Mức 1", "Mức 2", "Mức 3", "Mức 4", "Mức 5", "Số trả lời", "Điểm TB"])
    for k, c in enumerate(cau):
        r = k + 2
        ws.append([k + 1, c["noi_dung"], c["nhom"], "Có" if c["dao"] else "Không", *c["phan_bo_muc_1_den_5"], f"=SUM(E{r}:I{r})",
                   (f"=(5*E{r}+4*F{r}+3*G{r}+2*H{r}+1*I{r})/J{r}" if c["dao"] else f"=(1*E{r}+2*F{r}+3*G{r}+4*H{r}+5*I{r})/J{r}")])
    ws.append(["", "Điểm trung bình chung", "", "", "", "", "", "", "", "", "=AVERAGE(K2:K13)"])
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9D9D9")
    for col, wd in zip("ABCDEFGHIJK", [6, 40, 18, 9, 8, 8, 8, 8, 8, 11, 10]):
        ws.column_dimensions[col].width = wd
    for r in range(2, 15):
        ws[f"K{r}"].number_format = "0.00"
    ws.freeze_panes = "A2"
    bc = BarChart()
    bc.type = "col"
    bc.title = "Điểm trung bình từng câu hỏi"
    bc.y_axis.title = "Điểm trung bình (thang 1–5)"
    bc.x_axis.title = "Câu hỏi (STT)"
    bc.y_axis.scaling.min, bc.y_axis.scaling.max = 0, 5
    bc.add_data(Reference(ws, min_col=11, min_row=1, max_row=13), titles_from_data=True)
    bc.set_categories(Reference(ws, min_col=1, min_row=2, max_row=13))
    bc.legend = None
    bc.width, bc.height = 22, 10
    ws.add_chart(bc, "B17")
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.print_title_rows = "1:1"
    ws.print_area = "A1:W40"
    wb.save(O / "phan-tich-ket-qua-khao-sat.xlsx")

    xexp = {"skill": "phan-tich-ket-qua-khao-sat--xlsx", "output": "phan-tich-ket-qua-khao-sat.xlsx", "kind": "xlsx", "sheets": ["Du_lieu"],
            "charts": 1, "fit_to_width": True, "computed_approx": {"Du_lieu": {f"K{k + 2}": round(m, 2) for k, m in enumerate(tb_list)} | {"K14": round(tb_chung, 2)}},
            "must_contain": ["Điểm trung bình chung"], "allowed_derived": ["1.010", "1010"]}
    (S / "phan-tich-ket-qua-khao-sat--xlsx.input.json").write_text(json.dumps({"_ghi_chu": NOTE, **inp}, ensure_ascii=False, indent=1), encoding="utf-8")
    (S / "phan-tich-ket-qua-khao-sat--xlsx.expect.json").write_text(json.dumps(xexp, ensure_ascii=False, indent=1), encoding="utf-8")
    # ---- kỳ vọng ----
    tbq = {k + 1: vn(m) for k, m in enumerate(tb_list)}
    exp = {"skill": "phan-tich-ket-qua-khao-sat", "output": "phan-tich-ket-qua-khao-sat.docx", "kind": "docx", "nd30_format": True,
           "must_contain": ["BÁO CÁO", "I. THÔNG TIN CHUNG", "II. KẾT QUẢ CHI TIẾT", "III. NHẬN XÉT", "IV. HẠN CHẾ CỦA SỐ LIỆU", "V. ĐỀ XUẤT CẢI TIẾN",
                              "đã mã hóa ngược", "dưới 30", "TRƯỞNG PHÒNG KHẢO THÍ"],
           "must_not_contain": ["Khoa Luật 4,20", "Luật 4,20"],
           "blank_labels": ["Số:", "Hà Nội, ngày"],
           "allowed_derived": [str(n_valid), str(tong_loai), str(loai_thuc), "1.010"] + [str(v) for v in n_ans],
           "layout": {"page_number": True, "min_pages": 4,
                      "tables": [{"index": 1, "header_repeat": True, "cant_split": True},
                                 {"index": 2, "header_repeat": True, "cant_split": True, "data_rows": 12, "stt_continuous": True, "total_row_prefix": "Điểm trung bình chung"}],
                      "same_page": [["V. ĐỀ XUẤT CẢI TIẾN", "TRƯỞNG PHÒNG KHẢO THÍ"]],
                      "figures": [
                          {"no": 1, "table_index": 2, "label_col": 0, "value_col": 3, "label_fmt": "Câu {}", "min_points": 12},
                          {"no": 2, "min_points": 12}, {"no": 3, "min_points": 10}, {"no": 4, "min_points": 4}]},
           "traps": [{"id": "CAU_DAO", "mo_ta": "Câu 9 là câu đảo: phải mã hóa ngược trước khi tính điểm"},
                     {"id": "TONG_PHIEU_LOAI_LECH", "mo_ta": "Ba lý do loại cộng 170, thu về trừ hợp lệ là 168"},
                     {"id": "THIEU_TRA_LOI", "mo_ta": "Câu 7 có 12 phiếu bỏ trống: tính trên 830, không điền hộ"},
                     {"id": "DIEM_TB_LECH", "mo_ta": "Câu 3: điểm khai báo 4,05 khác điểm tính lại từ phân bố"},
                     {"id": "CAU_DOI_NOI_DUNG", "mo_ta": "Câu 5 đổi nội dung so với kỳ trước: không so sánh"},
                     {"id": "NHOM_NHO", "mo_ta": "Khoa Luật 24 phiếu (<30): không kết luận riêng, không vẽ"},
                     {"id": "KHOA_LECH_TB_CHUNG", "mo_ta": "TB chung theo khoa gia quyền khác TB chung 12 câu"},
                     {"id": "BIEU_DO_KHOP_BANG", "mo_ta": "Số liệu trong hình khớp bảng"}]}
    (S / "phan-tich-ket-qua-khao-sat.input.json").write_text(json.dumps({"_ghi_chu": NOTE, **inp}, ensure_ascii=False, indent=1), encoding="utf-8")
    (S / "phan-tich-ket-qua-khao-sat.expect.json").write_text(json.dumps(exp, ensure_ascii=False, indent=1), encoding="utf-8")
    q9_rev = tb_list[8]
    fnd = (f"- **CAU_DAO**: câu 9 điểm gốc {vn(q9_raw)} (xếp loại trung bình) nhưng là câu đảo; sau khi mã hóa ngược còn {vn(q9_rev)} (cần cải thiện). Nếu quên đảo, kết luận đổi chiều: đây là bẫy quan trọng nhất của bài.\n"
           f"- **TONG_PHIEU_LOAI_LECH**: ba lý do loại cộng {tong_loai}, thu về trừ hợp lệ là {loai_thuc} (lệch {tong_loai - loai_thuc}). Báo cáo giữ nguyên đầu vào, ghi lệch ở mục I và IV, không tự sửa.\n"
           f"- **THIEU_TRA_LOI**: câu 7 chỉ có {n_ans[6]} phiếu trả lời trên {n_valid}; điểm tính trên {n_ans[6]}, cột 'Số trả lời' hiển thị rõ, không điền hộ.\n"
           f"- **DIEM_TB_LECH**: câu 3 đầu vào ghi {vn(cau[2]['diem_tb_khai_bao'])}, tính lại {vn(tb_list[2])}; dùng giá trị tính lại và nêu ở mục hạn chế.\n"
           f"- **CAU_DOI_NOI_DUNG**: câu 5 không đưa vào so sánh kỳ trước (ô để trống, không ghi 0); hình 3 chỉ có {len(idx)} cặp cột.\n"
           f"- **NHOM_NHO**: khoa Luật {khoa_nho[0][1]} phiếu: không vẽ trong hình 4, không kết luận riêng (điểm khai báo {vn(khoa_nho[0][2])} sẽ gây hiểu nhầm là cao nhất).\n"
           f"- **KHOA_LECH_TB_CHUNG**: trung bình gia quyền theo khoa {vn(tb_khoa_gq)} so với trung bình 12 câu {vn(tb_chung)}; lệch {vn(abs(tb_khoa_gq - tb_chung))}. Nêu ở hạn chế, không tự quyết số nào đúng.\n"
           f"- **BIEU_DO_KHOP_BANG**: văn bản thay thế (alt) của hình 1 ghi lại điểm từng câu; bộ kiểm tra so với cột 'Điểm TB' của bảng 2.\n"
           f"- Cách ghi số: toàn văn bản dùng dấu phẩy thập phân (4,12) nhất quán; ví dụ trong SKILL.md đang viết '4.0'. Cần thống nhất trong skill.\n"
           f"- File Excel đi kèm có biểu đồ gốc (sửa được) và công thức điểm trung bình; file Word chèn ảnh PNG (không sửa số trong ảnh, phải sửa ở dữ liệu rồi vẽ lại).\n")
    (R / "phan-tich-ket-qua-khao-sat.findings.md").write_text("# Phát hiện khi chạy thử báo cáo có biểu đồ: phan-tich-ket-qua-khao-sat (dữ liệu giả)\n\n" + fnd, encoding="utf-8")
    return out


# =====================================================================
# 2. bao-cao-tien-do-de-tai
# =====================================================================
KP = [("Thù lao nhân công", 120_000_000, 70_000_000, 62_000_000), ("Nguyên vật liệu", 60_000_000, 30_000_000, 29_500_000),
      ("Thiết bị", 80_000_000, 40_000_000, 46_000_000), ("Hội thảo, công tác phí", 30_000_000, 15_000_000, 4_000_000),
      ("Quản lý phí", 10_000_000, 5_000_000, 5_000_000)]
ND = [("Tổng quan tài liệu", 100, 100, "Báo cáo tổng quan 28 trang"), ("Thu thập dữ liệu", 80, 70, "Bộ dữ liệu 1 200 mẫu"),
      ("Xây dựng mô hình", 50, 35, "Bản mô hình thử nghiệm v0.3"), ("Khảo sát thực nghiệm", 40, 100, None), ("Viết bài báo", 20, 10, "Bản thảo mục 1–2")]


def gen_tien_do():
    return {"co_quan_chu_quan": CQ, "co_quan_ban_hanh": TR + " – PHÒNG KHOA HỌC VÀ CÔNG NGHỆ", "dia_danh": "Hà Nội", "so_van_ban": None,
            "ten_de_tai": "Nghiên cứu mô hình dự báo kết quả học tập của sinh viên (tên mẫu)", "ma_so_de_tai": "ĐT.26.15 (mã mẫu)",
            "chu_nhiem": "TS. Nguyễn Văn Mẫu", "ky_bao_cao": "6 tháng đầu năm 2026 (01/01/2026–30/06/2026)",
            "noi_dung": [{"ten": n, "ke_hoach_luy_ke_phan_tram": kh, "thuc_hien_phan_tram_khai_bao": th, "minh_chung": mc} for n, kh, th, mc in ND],
            "kinh_phi": [{"khoan_muc": a, "du_toan_duoc_duyet": money(b), "da_cap": money(c), "da_su_dung": money(dd)} for a, b, c, dd in KP],
            "tong_da_cap_khai_bao": money(150_000_000), "tong_da_su_dung_khai_bao": money(146_500_000), "tong_du_toan_khai_bao": money(300_000_000),
            "khong_so_san": "", "noi_nhan": ["Phòng Khoa học và Công nghệ", "Lưu: VT, KHCN"], "nguoi_ky": {"chuc_danh": "Chủ nhiệm đề tài", "ho_ten": "TS. Nguyễn Văn Mẫu"}}


def build_tien_do():
    inp = gen_tien_do()
    tot_dt = sum(x[1] for x in KP)
    tot_cap = sum(x[2] for x in KP)
    tot_dung = sum(x[3] for x in KP)
    # Hình 1
    fig, ax = plt.subplots(figsize=(6.4, 3.6))
    xs = list(range(len(KP)))
    w = 0.27
    series = [("Dự toán được duyệt", [x[1] for x in KP]), ("Đã cấp", [x[2] for x in KP]), ("Đã sử dụng", [x[3] for x in KP])]
    for s, (nm, vals) in enumerate(series):
        bs = ax.bar([x + (s - 1) * w for x in xs], [v / 1e6 for v in vals], w, color=MAU[s], hatch=HATCH[s], edgecolor="#1F3864", label=nm)
        for b, v in zip(bs, vals):
            ax.text(b.get_x() + b.get_width() / 2, b.get_height() + 1.5, f"{v / 1e6:.1f}".replace(".", ","), ha="center", fontsize=7)
    ax.set_xticks(xs)
    ax.set_xticklabels([x[0].replace(", ", ",\n").replace(" nhân ", " nhân\n") for x in KP], fontsize=8.5)
    ax.set_ylabel("Triệu đồng")
    ax.legend(loc="upper center", bbox_to_anchor=(0.5, -0.2), ncol=3, frameon=False, fontsize=8.5)
    p1 = save(fig, "tien-do-hinh1")
    # Hình 2: % hoàn thành (nội dung thiếu minh chứng không vẽ cột thực hiện)
    fig, ax = plt.subplots(figsize=(6.4, 3.0))
    names = [n for n, *_ in ND][::-1]
    kh = [k for _, k, *_ in ND][::-1]
    th = [t if mc else None for _, _, t, mc in ND][::-1]
    ys = list(range(len(ND)))
    ax.barh([y + 0.19 for y in ys], kh, 0.36, color=MAU[2], hatch="///", edgecolor="#1F3864", label="Kế hoạch lũy kế đến 30/06/2026")
    for y, t in zip(ys, th):
        if t is not None:
            ax.barh(y - 0.19, t, 0.36, color=MAU[0], edgecolor="#1F3864", label="Thực hiện (có minh chứng)" if y == ys[0] else None)
            ax.text(t + 1.5, y - 0.19, f"{t}%", va="center", fontsize=8)
        else:
            ax.text(2, y - 0.19, "chưa có minh chứng", va="center", fontsize=8, style="italic")
    for y, k in zip(ys, kh):
        ax.text(k + 1.5, y + 0.19, f"{k}%", va="center", fontsize=8)
    ax.set_yticks(ys)
    ax.set_yticklabels(names, fontsize=8.5)
    ax.set_xlim(0, 115)
    ax.set_xlabel("Tỷ lệ hoàn thành (%)")
    ax.grid(axis="y", visible=False)
    h, l_ = ax.get_legend_handles_labels()
    ax.legend(h, l_, loc="upper center", bbox_to_anchor=(0.5, -0.22), ncol=2, frameon=False, fontsize=8.5)
    p2 = save(fig, "tien-do-hinh2")

    d = new_doc()
    page_number(d)
    head(d, inp, "BC")
    para(d, "BÁO CÁO", bold=True, size=14, align=C, after=0)
    para(d, "TIẾN ĐỘ THỰC HIỆN ĐỀ TÀI NGHIÊN CỨU KHOA HỌC", bold=True, align=C, after=0)
    para(d, "6 tháng đầu năm 2026", bold=True, align=C, after=8)
    for lab, key in (("Tên đề tài", "ten_de_tai"), ("Mã số", "ma_so_de_tai"), ("Chủ nhiệm", "chu_nhiem"), ("Kỳ báo cáo", "ky_bao_cao")):
        para(d, f"{lab}: {inp[key]}", first_indent=1.0, align=J, after=6)
    para(d, "", after=4)
    para(d, "I. NỘI DUNG ĐÃ THỰC HIỆN TRONG KỲ", bold=True, after=4, keep=True)
    para(d, "Tỷ lệ hoàn thành chỉ ghi khi có minh chứng kèm theo; nội dung chưa có minh chứng để trống.", first_indent=1.0, align=J)
    add_figure(d, p2, 1, "Tỷ lệ hoàn thành từng nội dung nghiên cứu so với kế hoạch lũy kế", "Báo cáo của chủ nhiệm đề tài; kế hoạch theo thuyết minh đã duyệt",
               alt_of("Biểu đồ thanh ngang kế hoạch và thực hiện", [(n, f"{k}->{t if mc else 'chưa có minh chứng'}") for n, k, t, mc in ND], "% hoàn thành"))
    rows = []
    for k, (n, kh_, th_, mc) in enumerate(ND, 1):
        rows.append([str(k), n, f"{kh_}%", (f"{th_}%" if mc else ""), mc or "Chưa có minh chứng"])
    para(d, "Bảng 1. Đánh giá từng nội dung nghiên cứu", bold=True, size=13, align=C, after=2, keep=True)
    grid_k(d, ["STT", "Nội dung", "Kế hoạch", "Thực hiện", "Minh chứng"], rows, [1.4, 4.6, 2.2, 2.2, 5.6], aligns=[C, L, C, C, L], size=11)
    para(d, "", after=6)
    para(d, "II. TÌNH HÌNH SỬ DỤNG KINH PHÍ", bold=True, after=4, keep=True)
    add_figure(d, p1, 2, "Dự toán, kinh phí đã cấp và đã sử dụng theo khoản mục", "Sổ sách kế toán của đề tài do chủ nhiệm cung cấp",
               alt_of("Biểu đồ cột ghép theo khoản mục", [(a, f"{b / 1e6:.1f}/{c / 1e6:.1f}/{dd / 1e6:.1f}") for a, b, c, dd in KP], "triệu đồng: dự toán/đã cấp/đã sử dụng"))
    rows = []
    for k, (a, b, c, dd) in enumerate(KP, 1):
        pct = vn(100 * dd / c, 1) + "%"
        rows.append([str(k), a, money(b), money(c), money(dd), pct])
    para(d, "Bảng 2. Kinh phí theo khoản mục (đồng)", bold=True, size=13, align=C, after=2, keep=True)
    grid_k(d, ["STT", "Khoản mục", "Dự toán", "Đã cấp", "Đã sử dụng", "% đã dùng / đã cấp"], rows, [1.4, 4.2, 2.6, 2.6, 2.6, 2.6],
         total=["", "Cộng", money(tot_dt), "", money(tot_dung), ""], aligns=[C, L, "R" and WD_ALIGN_PARAGRAPH.RIGHT, WD_ALIGN_PARAGRAPH.RIGHT, WD_ALIGN_PARAGRAPH.RIGHT, C], size=11)
    para(d, f"Lưu ý: tổng đã cấp theo đầu vào là {inp['tong_da_cap_khai_bao']} đồng, nhưng cộng các khoản mục được {money(tot_cap)} đồng; chưa ghi tổng đã cấp vào bảng cho đến khi kế toán xác nhận. "
            f"Khoản mục Thiết bị đã sử dụng {money(KP[2][3])} đồng, cao hơn đã cấp {money(KP[2][2])} đồng (đạt {vn(100 * KP[2][3] / KP[2][2], 1)}%): cần giải trình nguồn bù. "
            f"Khoản mục Nguyên vật liệu đạt {vn(100 * KP[1][3] / KP[1][2], 1)}% (trên 95%) và khoản mục Hội thảo, công tác phí đạt {vn(100 * KP[3][3] / KP[3][2], 1)}% (dưới 50%): đều cần giải trình.",
         italic=True, size=13, first_indent=1.0, align=J, after=8)
    para(d, "III. KHÓ KHĂN, VƯỚNG MẮC", bold=True, after=4, keep=True)
    para(d, DOT * 3, first_indent=1.0, after=8)
    para(d, "IV. ĐỀ XUẤT, KIẾN NGHỊ", bold=True, after=4, keep=True)
    para(d, DOT * 3, first_indent=1.0, after=8)
    para(d, "V. KẾ HOẠCH KỲ TIẾP THEO", bold=True, after=4, keep=True)
    para(d, DOT * 3, first_indent=1.0, after=8, keep=True)
    keep_together(d, 2)
    sign(d, inp["noi_nhan"], inp["nguoi_ky"]["chuc_danh"], inp["nguoi_ky"]["ho_ten"])
    out = O / "bao-cao-tien-do-de-tai.docx"
    d.save(out)
    exp = {"skill": "bao-cao-tien-do-de-tai", "output": "bao-cao-tien-do-de-tai.docx", "kind": "docx", "nd30_format": True,
           "must_contain": ["I. NỘI DUNG ĐÃ THỰC HIỆN TRONG KỲ", "II. TÌNH HÌNH SỬ DỤNG KINH PHÍ", "III. KHÓ KHĂN, VƯỚNG MẮC", "IV. ĐỀ XUẤT, KIẾN NGHỊ", "V. KẾ HOẠCH KỲ TIẾP THEO",
                              "Chưa có minh chứng", "cần giải trình nguồn bù", "CHỦ NHIỆM ĐỀ TÀI"],
           "must_not_contain": ["150.000.000 đồng, nhưng cộng các khoản mục được 150"],
           "blank_labels": ["Số:", "Hà Nội, ngày"],
           "allowed_derived": [money(tot_dt), money(tot_cap), money(tot_dung), "115,0"],
           "layout": {"page_number": True, "min_pages": 3,
                      "tables": [{"index": 1, "header_repeat": True, "cant_split": True, "data_rows": 5, "stt_continuous": True},
                                 {"index": 2, "header_repeat": True, "cant_split": True, "data_rows": 5, "stt_continuous": True, "total_row_prefix": "Cộng"}],
                      "same_page": [["V. KẾ HOẠCH KỲ TIẾP THEO", "CHỦ NHIỆM ĐỀ TÀI"]],
                      "figures": [{"no": 1, "min_points": 5}, {"no": 2, "min_points": 5}]},
           "traps": [{"id": "DA_DUNG_VUOT_DA_CAP", "mo_ta": "Thiết bị: dùng 46 triệu > cấp 40 triệu"}, {"id": "TONG_DA_CAP_LECH", "mo_ta": "Tổng đã cấp 150 triệu, cộng khoản mục 160 triệu"},
                     {"id": "TY_LE_NGOAI_NGUONG", "mo_ta": "Nguyên vật liệu 98,3%, Hội thảo 26,7%: cần giải trình"},
                     {"id": "THIEU_MINH_CHUNG", "mo_ta": "Khảo sát thực nghiệm khai 100% nhưng không có minh chứng: để trống, không vẽ"},
                     {"id": "MUC_CHUA_CO_DU_LIEU", "mo_ta": "Khó khăn, đề xuất, kế hoạch kỳ sau chưa có dữ liệu: chừa trống"},
                     {"id": "BIEU_DO_KHOP_BANG", "mo_ta": "Hình 2 khớp bảng 2"}]}
    (S / "bao-cao-tien-do-de-tai.input.json").write_text(json.dumps({"_ghi_chu": NOTE, **inp}, ensure_ascii=False, indent=1), encoding="utf-8")
    (S / "bao-cao-tien-do-de-tai.expect.json").write_text(json.dumps(exp, ensure_ascii=False, indent=1), encoding="utf-8")
    fnd = (f"- **DA_DUNG_VUOT_DA_CAP**: Thiết bị đã dùng {money(KP[2][3])} > đã cấp {money(KP[2][2])} ({vn(100 * KP[2][3] / KP[2][2], 1)}%). Báo cáo giữ số, nêu cần giải trình nguồn bù; hình 2 vẽ đúng cột vượt.\n"
           f"- **TONG_DA_CAP_LECH**: đầu vào khai tổng {inp['tong_da_cap_khai_bao']}, cộng khoản mục {money(tot_cap)}. Ô tổng đã cấp để trống, nêu lệch trong chú thích.\n"
           f"- **TY_LE_NGOAI_NGUONG**: Nguyên vật liệu {vn(100 * KP[1][3] / KP[1][2], 1)}% (>95%), Hội thảo {vn(100 * KP[3][3] / KP[3][2], 1)}% (<50%): ghi cần giải trình theo ngưỡng của skill.\n"
           f"- **THIEU_MINH_CHUNG**: nội dung 'Khảo sát thực nghiệm' khai 100% nhưng minh chứng trống. Skill cấm ghi kết quả không có minh chứng: ô Thực hiện để trống, trong hình thay cột bằng chữ 'chưa có minh chứng' (không vẽ 0%).\n"
           f"- **MUC_CHUA_CO_DU_LIEU**: mục III, IV, V không có dữ liệu đầu vào; để dòng chấm, không bịa nội dung.\n"
           f"- **BIEU_DO_KHOP_BANG**: alt của hình ghi số liệu từng khoản mục/nội dung; bộ kiểm tra đối chiếu với bảng.\n")
    (R / "bao-cao-tien-do-de-tai.findings.md").write_text("# Phát hiện khi chạy thử báo cáo có biểu đồ: bao-cao-tien-do-de-tai (dữ liệu giả)\n\n" + fnd, encoding="utf-8")
    return out


# =====================================================================
# 3. bao-cao-khcn-nam
# =====================================================================
CB = {"WoS/Scopus": [18, 24, 31, 22], "Tạp chí trong nước": [46, 51, 49, 30], "Kỷ yếu hội thảo": [None, 28, 25, 15]}
NAMS = ["2023", "2024", "2025", "2026 (đến 30/9)"]
DT = [("Nhà nước", 2, 3200), ("Bộ", 5, 4100), ("Tỉnh", 3, 1500), ("Trường", 22, 2640)]


def build_khcn():
    inp = {"co_quan_chu_quan": CQ, "co_quan_ban_hanh": TR, "dia_danh": "Hà Nội", "so_van_ban": None, "nam_bao_cao": "2026",
           "ghi_chu_thoi_diem": "Số liệu năm 2026 tính đến 30/9/2026",
           "cong_bo_khoa_hoc": {"nam": NAMS, **{k: v for k, v in CB.items()}}, "tong_cong_bo_2025_khai_bao": 107,
           "de_tai_cac_cap": [{"cap": c, "so_de_tai": n, "kinh_phi_trieu_dong": money(k)} for c, n, k in DT],
           "tong_kinh_phi_khai_bao_trieu_dong": money(11440), "noi_nhan": ["Bộ Giáo dục và Đào tạo", "Lưu: VT, KHCN"],
           "nguoi_ky": {"chuc_danh": "Hiệu trưởng", "ho_ten": None}}
    tot = []
    for i in range(4):
        vals = [CB[k][i] for k in CB]
        tot.append(None if any(v is None for v in vals) else sum(vals))
    fig, ax = plt.subplots(figsize=(6.4, 3.6))
    bottom = [0] * 4
    for s, (k, vals) in enumerate(CB.items()):
        v = [x or 0 for x in vals]
        bs = ax.bar(NAMS, v, 0.55, bottom=bottom, color=MAU[s], hatch=HATCH[s], edgecolor="#1F3864", label=k)
        for b, x, bt in zip(bs, vals, bottom):
            if x is None:
                ax.text(b.get_x() + b.get_width() / 2, bt + 3, "n/a", ha="center", fontsize=8, style="italic")
            elif x >= 8:
                ax.text(b.get_x() + b.get_width() / 2, bt + x / 2, str(x), ha="center", va="center", fontsize=8.5, color="white" if s == 0 else "black")
        bottom = [a + b for a, b in zip(bottom, v)]
    ax.set_ylabel("Số công trình công bố")
    ax.legend(loc="upper center", bbox_to_anchor=(0.5, -0.16), ncol=3, frameon=False, fontsize=8.5)
    ax.tick_params(axis="x", labelsize=8.5)
    p1 = save(fig, "khcn-hinh1")
    fig, ax = plt.subplots(figsize=(6.4, 2.8))
    bs = ax.bar([c for c, _, _ in DT], [k for _, _, k in DT], 0.5, color=MAU[1], edgecolor="#1F3864")
    for b, (c, n, k) in zip(bs, DT):
        ax.text(b.get_x() + b.get_width() / 2, k + 60, f"{money(k)}\n({n} đề tài)", ha="center", fontsize=8.5)
    ax.set_ylim(0, 5200)
    ax.set_ylabel("Triệu đồng")
    ax.yaxis.set_major_formatter(matplotlib.ticker.FuncFormatter(lambda v, _: money(int(v))))
    p2 = save(fig, "khcn-hinh2")
    d = new_doc()
    page_number(d)
    head(d, inp, "BC")
    para(d, "BÁO CÁO", bold=True, size=14, align=C, after=0)
    para(d, "HOẠT ĐỘNG KHOA HỌC VÀ CÔNG NGHỆ NĂM 2026", bold=True, align=C, after=0)
    para(d, "(Số liệu năm 2026 tính đến 30/9/2026)", italic=True, align=C, after=8)
    para(d, "I. KẾT QUẢ THỰC HIỆN", bold=True, after=4, keep=True)
    para(d, "1. Công bố khoa học", bold=True, after=4, keep=True)
    para(d, "Số liệu năm 2026 mới tính đến 30/9/2026 nên chưa so sánh trực tiếp với cả năm của các năm trước. Số kỷ yếu hội thảo năm 2023 chưa có số liệu (ghi n/a, không ghi 0), do đó chưa có tổng năm 2023.",
         first_indent=1.0, align=J)
    add_figure(d, p1, 1, "Số công trình khoa học công bố theo loại, giai đoạn 2023–2026", "Số liệu do Phòng Khoa học và Công nghệ tổng hợp",
               alt_of("Biểu đồ cột xếp chồng theo năm", [(f"{n} - {k}", ("n/a" if CB[k][i] is None else CB[k][i])) for i, n in enumerate(NAMS) for k in CB], "công trình"),
               note="Năm 2026 tính đến 30/9/2026")
    rows = []
    for i, n in enumerate(NAMS):
        rows.append([str(i + 1), n] + [("" if CB[k][i] is None else str(CB[k][i])) for k in CB] + [("" if tot[i] is None else str(tot[i]))])
    para(d, "Bảng 1. Công bố khoa học theo năm", bold=True, size=13, align=C, after=2, keep=True)
    grid_k(d, ["STT", "Năm", "WoS/Scopus", "Trong nước", "Kỷ yếu", "Cộng"], rows, [1.4, 3.8, 2.6, 2.6, 2.6, 3.0], aligns=[C, L, C, C, C, C], size=11)
    para(d, f"Lưu ý: đầu vào khai tổng công bố năm 2025 là {inp['tong_cong_bo_2025_khai_bao']}, cộng ba loại là {tot[2]}; báo cáo ghi theo tổng các loại và đề nghị đơn vị xác nhận.",
         italic=True, size=13, first_indent=1.0, align=J, after=8)
    para(d, "2. Đề tài các cấp", bold=True, after=4, keep=True)
    add_figure(d, p2, 2, "Kinh phí đề tài khoa học và công nghệ theo cấp quản lý (triệu đồng)", "Số liệu do Phòng Khoa học và Công nghệ tổng hợp",
               alt_of("Biểu đồ cột kinh phí theo cấp", [(c, k) for c, _, k in DT], "triệu đồng"))
    rows = [[str(i + 1), c, str(n), money(k)] for i, (c, n, k) in enumerate(DT)]
    tong = sum(k for _, _, k in DT)
    para(d, "Bảng 2. Đề tài theo cấp quản lý", bold=True, size=13, align=C, after=2, keep=True)
    grid_k(d, ["STT", "Cấp", "Số đề tài", "Kinh phí (triệu đồng)"], rows, [1.4, 5.0, 3.6, 6.0], total=["", "Cộng", str(sum(n for _, n, _ in DT)), money(tong)],
         aligns=[C, L, C, WD_ALIGN_PARAGRAPH.RIGHT], size=11)
    para(d, "II. ĐÁNH GIÁ CHUNG", bold=True, after=4, keep=True)
    para(d, DOT * 3, first_indent=1.0, after=8)
    para(d, "III. PHƯƠNG HƯỚNG NĂM TIẾP THEO", bold=True, after=4, keep=True)
    para(d, DOT * 3, first_indent=1.0, after=8, keep=True)
    keep_together(d, 2)
    sign(d, inp["noi_nhan"], inp["nguoi_ky"]["chuc_danh"], inp["nguoi_ky"]["ho_ten"])
    out = O / "bao-cao-khcn-nam.docx"
    d.save(out)
    exp = {"skill": "bao-cao-khcn-nam", "output": "bao-cao-khcn-nam.docx", "kind": "docx", "nd30_format": True,
           "must_contain": ["I. KẾT QUẢ THỰC HIỆN", "II. ĐÁNH GIÁ CHUNG", "III. PHƯƠNG HƯỚNG NĂM TIẾP THEO", "đến 30/9/2026", "n/a", "HIỆU TRƯỞNG"],
           "must_not_contain": [], "blank_labels": ["Số:", "Hà Nội, ngày"], "allowed_derived": [money(tong), "103", "105"],
           "layout": {"page_number": True, "min_pages": 2, "same_page": [["III. PHƯƠNG HƯỚNG NĂM TIẾP THEO", "HIỆU TRƯỞNG"]],
                      "tables": [{"index": 1, "header_repeat": True, "cant_split": True, "data_rows": 4, "stt_continuous": True},
                                 {"index": 2, "header_repeat": True, "cant_split": True, "data_rows": 4, "stt_continuous": True, "total_row_prefix": "Cộng"}],
                      "figures": [{"no": 1, "min_points": 12}, {"no": 2, "table_index": 2, "label_col": 1, "value_col": 3, "min_points": 4}]},
           "traps": [{"id": "NAM_CHUA_DU", "mo_ta": "2026 chỉ đến 30/9: không so như cả năm"}, {"id": "THIEU_SO_LIEU_KHAC_0", "mo_ta": "Kỷ yếu 2023 không có số: ghi n/a, không ghi 0, không cộng tổng"},
                     {"id": "TONG_2025_LECH", "mo_ta": "Khai 107, cộng 105"}, {"id": "BIEU_DO_KHOP_BANG", "mo_ta": "Hình 2 khớp bảng 2"},
                     {"id": "MUC_CHUA_CO_DU_LIEU", "mo_ta": "Đánh giá chung, phương hướng chưa có dữ liệu: chừa trống"}]}
    (S / "bao-cao-khcn-nam.input.json").write_text(json.dumps({"_ghi_chu": NOTE, **inp}, ensure_ascii=False, indent=1), encoding="utf-8")
    (S / "bao-cao-khcn-nam.expect.json").write_text(json.dumps(exp, ensure_ascii=False, indent=1), encoding="utf-8")
    fnd = ("- **NAM_CHUA_DU**: số liệu 2026 chỉ đến 30/9; tên cột, ghi chú dưới hình, chú thích văn bản đều nói rõ. Không nhận xét 'giảm so với 2025' vì so một nửa năm với cả năm.\n"
           "- **THIEU_SO_LIEU_KHAC_0**: kỷ yếu 2023 không có số; bảng để ô trống, hình ghi 'n/a', ô tổng 2023 để trống (không cộng thiếu).\n"
           f"- **TONG_2025_LECH**: khai {inp['tong_cong_bo_2025_khai_bao']}, cộng {tot[2]}; ghi chú nêu lệch, bảng dùng tổng các loại.\n"
           "- **BIEU_DO_KHOP_BANG**: alt của hình 2 ghi kinh phí từng cấp, đối chiếu bảng 2.\n"
           "- **MUC_CHUA_CO_DU_LIEU**: mục II, III không có dữ liệu đầu vào; dòng chấm.\n")
    (R / "bao-cao-khcn-nam.findings.md").write_text("# Phát hiện khi chạy thử báo cáo có biểu đồ: bao-cao-khcn-nam (dữ liệu giả)\n\n" + fnd, encoding="utf-8")
    return out


if __name__ == "__main__":
    for f in (build_khao_sat, build_tien_do, build_khcn):
        print("đã tạo", f())
