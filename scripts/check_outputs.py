"""Kiểm tra file đầu ra của skill theo các quy tắc trong quy cách đầu ra.

Không phụ thuộc cách file được tạo: dùng được cho file do AI nào tạo ra từ skill.
Chạy: python -X utf8 scripts/check_outputs.py [--samples tests/samples] [--outputs tests/outputs] [--results tests/results]

Với mỗi cặp <skill>.input.json / <skill>.expect.json kiểm tra:
  1. File tồn tại, mở được bằng thư viện thật (không đổi đuôi giả).
  2. Không có chuỗi bị cấm (nhãn chờ ký, checklist, "chưa cung cấp", số 0 thay dữ liệu...).
  3. Không bịa số liệu: mọi ngày và số từ 3 chữ số trong file phải có trong đầu vào.
  4. Có/không có các chuỗi bắt buộc; trường thiếu dữ liệu được chừa chỗ điền.
  5. Với Excel: tính lại công thức bằng LibreOffice rồi so các ô đã khai báo.
  6. Mỗi "bẫy" trong expect có ghi nhận trong tests/results/<skill>.findings.md.
"""
import argparse
import json
import re
import shutil
import subprocess
import sys
import tempfile
from datetime import datetime
from pathlib import Path

GLOBAL_FORBIDDEN = ["CHỜ KÝ", "CHƯA CUNG CẤP", "chưa cung cấp", "cần xác minh", "CẦN XÁC MINH", "(đã ký)",
                    "DỰ THẢO – CHỜ KIỂM DUYỆT", "checklist", "Checklist", "phụ lục kiểm tra", "Phụ lục kiểm tra",
                    "☐", "[ ]", "TODO", "Lorem", "nhật ký AI"]


def docx_text(path):
    from docx import Document
    d = Document(path)
    out = [p.text for p in d.paragraphs]

    def walk(tables):
        for t in tables:
            for row in t.rows:
                for c in row.cells:
                    out.extend(p.text for p in c.paragraphs)
                    walk(c.tables)
    walk(d.tables)
    return "\n".join(x for x in out if x.strip())


def nd30_format_errors(path):
    """Kiểm tra thuộc tính trình bày thật của file Word theo Phụ lục I Nghị định 30/2020/NĐ-CP (văn bản hành chính)."""
    from docx import Document
    d = Document(path)
    errs = []
    s = d.sections[0]
    mm = lambda v: round(v / 36000, 1)  # EMU -> mm
    w, h = mm(s.page_width), mm(s.page_height)
    if not (abs(w - 210) <= 1 and abs(h - 297) <= 1):
        errs.append(f"khổ giấy không phải A4 ({w}x{h} mm)")
    for name, v, lo, hi in (("lề trên", mm(s.top_margin), 20, 25), ("lề dưới", mm(s.bottom_margin), 20, 25),
                            ("lề trái", mm(s.left_margin), 30, 35), ("lề phải", mm(s.right_margin), 15, 20)):
        if not (lo - 0.5 <= v <= hi + 0.5):
            errs.append(f"{name} {v} mm ngoài khoảng {lo}–{hi} mm")
    fonts, sizes = set(), set()

    def runs(doc_or_cell):
        for p in doc_or_cell.paragraphs:
            for r in p.runs:
                yield r
        for tb in getattr(doc_or_cell, "tables", []):
            for row in tb.rows:
                for c in row.cells:
                    yield from runs(c)
    normal = d.styles["Normal"].font
    for r in runs(d):
        if not r.text.strip():
            continue
        fonts.add(r.font.name or normal.name)
        sz = r.font.size or normal.size
        sizes.add(sz.pt if sz else None)
    if fonts - {"Times New Roman"}:
        errs.append(f"phông không phải Times New Roman: {sorted(f for f in fonts if f)}")
    bad = sorted(x for x in sizes if x is None or not (11 <= x <= 14))
    if bad:
        errs.append(f"cỡ chữ ngoài 11–14 pt: {bad}")
    return errs


def xlsx_load(path):
    from openpyxl import load_workbook
    return load_workbook(path)


def xlsx_text_and_numbers(wb):
    texts, nums = [], []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if v is None or (isinstance(v, str) and v.startswith("=")):
                    continue
                if isinstance(v, datetime):
                    texts.append(v.strftime("%d/%m/%Y"))
                elif isinstance(v, (int, float)):
                    nums.append(float(v))
                else:
                    texts.append(str(v))
    return "\n".join(texts), nums


def recalc(path):
    """Tính lại công thức bằng LibreOffice; trả về workbook data_only hoặc None."""
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        return None
    from openpyxl import load_workbook
    tmp = Path(tempfile.mkdtemp())
    src = tmp / "in" / path.name
    src.parent.mkdir()
    shutil.copy(path, src)
    outdir = tmp / "out"
    subprocess.run([soffice, "--headless", "--convert-to", "xlsx", "--outdir", str(outdir), str(src)],
                   capture_output=True, timeout=180)
    res = outdir / path.name
    return load_workbook(res, data_only=True) if res.exists() else None


def allowed_tokens(input_text):
    dates = set(re.findall(r"\d{1,2}/\d{1,2}/\d{4}", input_text))
    nums = set(re.findall(r"\d{3,}", input_text))
    return dates, nums


def pdf_pages(path):
    """Chuyển Word sang PDF bằng LibreOffice; trả về danh sách văn bản từng trang (hoặc None nếu không có công cụ)."""
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice or not shutil.which("pdftotext"):
        return None
    tmp = Path(tempfile.mkdtemp())
    src = tmp / path.name
    shutil.copy(path, src)
    subprocess.run([soffice, "--headless", "--convert-to", "pdf", "--outdir", str(tmp), str(src)],
                   capture_output=True, timeout=240)
    pdf = tmp / (path.stem + ".pdf")
    if not pdf.exists():
        return None
    out = subprocess.run(["pdftotext", "-layout", str(pdf), "-"], capture_output=True, text=True).stdout
    return out.split("\f")[:-1] if out.endswith("\f") else out.split("\f")


def layout_errors(path, expect, text, log):
    """Kiểm tra file nhiều dữ liệu/nhiều trang: số trang, số trang ở đầu trang, bảng, token duy nhất, ràng buộc cùng trang."""
    from docx import Document
    lay = expect.get("layout", {})
    errs = []
    d = Document(path)
    if lay.get("page_number"):  # NĐ 30/2020: đánh số trang từ trang 2, đặt giữa lề trên
        sect = d.sections[0]
        xml = sect._sectPr.xml
        hdr = "".join(p._p.xml for p in sect.header.paragraphs)
        if "w:titlePg" not in xml:
            errs.append("không bật 'trang đầu khác' nên số trang sẽ hiện cả ở trang 1")
        if "PAGE" not in hdr:
            errs.append("đầu trang không có trường số trang (PAGE)")
        elif 'w:jc w:val="center"' not in hdr:
            errs.append("số trang không căn giữa")
    for spec in lay.get("tables", []):
        if spec["index"] >= len(d.tables):
            errs.append(f"không có bảng thứ {spec['index']}")
            continue
        tb = d.tables[spec["index"]]
        rows = tb.rows
        if spec.get("header_repeat") and "w:tblHeader" not in rows[0]._tr.xml:
            errs.append(f"bảng {spec['index']}: dòng tiêu đề không lặp lại ở mỗi trang")
        if spec.get("cant_split") and not all("w:cantSplit" in r._tr.xml for r in rows[1:]):
            errs.append(f"bảng {spec['index']}: có dòng bị phép ngắt giữa hai trang")
        body = [[c.text.strip() for c in r.cells] for r in rows[1:]]
        if not body:
            errs.append(f"bảng {spec['index']}: không có dòng dữ liệu")
            continue
        n_total = len(spec.get("total_row_prefix", "")) and 1 or 0
        data = body[:-1] if n_total else body
        if "data_rows" in spec and len(data) != spec["data_rows"]:
            errs.append(f"bảng {spec['index']}: có {len(data)} dòng dữ liệu, kỳ vọng {spec['data_rows']}")
        if spec.get("stt_continuous"):
            stt = [r[0] for r in data]
            if stt != [str(k) for k in range(1, len(data) + 1)]:
                errs.append(f"bảng {spec['index']}: cột STT không liên tục từ 1")
        if n_total and not body[-1][0].startswith(spec["total_row_prefix"]) and spec["total_row_prefix"] not in " ".join(body[-1]):
            errs.append(f"bảng {spec['index']}: thiếu dòng tổng cộng")
    for tok in lay.get("unique_tokens", []):
        c = text.count(tok)
        if c != 1:
            errs.append(f"{tok!r} xuất hiện {c} lần, kỳ vọng đúng 1 lần")
    for tok in lay.get("absent_tokens", []):
        if tok in text:
            errs.append(f"{tok!r} không được có trong file")
    need_pdf = lay.get("min_pages") or lay.get("same_page")
    if need_pdf:
        pages = pdf_pages(path)
        if pages is None:
            log.append("  (bỏ qua kiểm số trang: không có LibreOffice/pdftotext)")
        else:
            if lay.get("min_pages") and len(pages) < lay["min_pages"]:
                errs.append(f"chỉ {len(pages)} trang, kỳ vọng ít nhất {lay['min_pages']}")
            log.append(f"  số trang thực tế: {len(pages)}")
            for a, b in lay.get("same_page", []):
                if not any(a in p and b in p for p in pages):
                    errs.append(f"{a!r} và {b!r} không cùng một trang (khối ký bị tách khỏi nội dung)")
            if lay.get("page_number") and len(pages) > 1:
                for k, p in enumerate(pages[1:], 2):
                    first = [ln.strip() for ln in p.split("\n") if ln.strip()][:1]
                    if not first or first[0] != str(k):
                        errs.append(f"trang {k} không có số trang ở đầu trang (thấy {first})")
                        break
                first1 = [ln.strip() for ln in pages[0].split("\n") if ln.strip()][:1]
                if first1 and first1[0] == "1":
                    errs.append("trang 1 không được hiện số trang")
    return errs


def check_one(expect, samples, outputs, results, log):
    errs = []
    name = expect["skill"]
    inp = (samples / f"{name}.input.json").read_text(encoding="utf-8")
    path = outputs / expect["output"]
    if not path.is_file():
        return [f"không thấy file {path.name}"]
    magic = path.read_bytes()[:4]
    if magic != b"PK\x03\x04":
        return ["không phải gói OOXML thật (đuôi file giả?)"]
    dates_ok, nums_ok = allowed_tokens(inp)
    for spec in expect.get("layout", {}).get("tables", []):  # số thứ tự dòng (STT) là số do bảng tự sinh
        nums_ok |= {str(k) for k in range(1, spec.get("data_rows", 0) + 1) if k >= 100}
    for s in expect.get("allowed_derived", []):  # số suy ra bằng phép tính từ đầu vào (tổng cộng...)
        nums_ok |= set(re.findall(r"\d{3,}", str(s)))
    wb_values = None
    try:
        if expect["kind"] == "docx":
            docx_text(path)
        else:
            xlsx_load(path)
    except Exception as e:  # file hỏng: báo lỗi, không làm sập bộ kiểm tra
        return [f"không mở được bằng thư viện ({type(e).__name__})"]
    if expect["kind"] == "docx":
        text = docx_text(path)
        nums = []
        if expect.get("nd30_format"):
            errs.extend(nd30_format_errors(path))
        if expect.get("layout"):
            errs.extend(layout_errors(path, expect, text, log))
    else:
        wb = xlsx_load(path)
        text, nums = xlsx_text_and_numbers(wb)
        sheets = [ws.title for ws in wb.worksheets]
        for s in expect.get("sheets", []):
            if s not in sheets:
                errs.append(f"thiếu sheet {s}")
        for s in expect.get("must_not_have_sheets", []):
            if s in sheets:
                errs.append(f"không được có sheet {s}")
        wb_values = recalc(path)
        if wb_values is None:
            log.append("  (bỏ qua kiểm công thức: không có LibreOffice)")
        else:
            for sheet, cells in expect.get("computed", {}).items():
                for ref, want in cells.items():
                    got = wb_values[sheet][ref].value
                    got = "" if got is None else got
                    if got != want:
                        errs.append(f"{sheet}!{ref}: được {got!r}, kỳ vọng {want!r}")
            for ws in wb_values.worksheets:  # không để lỗi công thức
                for row in ws.iter_rows():
                    for c in row:
                        if isinstance(c.value, str) and c.value.startswith("#"):
                            errs.append(f"lỗi công thức {ws.title}!{c.coordinate}: {c.value}")
    for bad in GLOBAL_FORBIDDEN + expect.get("must_not_contain", []):
        if bad in text:
            errs.append(f"có chuỗi bị cấm: {bad!r}")
    for need in expect.get("must_contain", []):
        if need not in text:
            errs.append(f"thiếu chuỗi bắt buộc: {need!r}")
    for d in set(re.findall(r"\d{1,2}/\d{1,2}/\d{4}", text)):
        if d not in dates_ok:
            errs.append(f"ngày không có trong đầu vào (nghi bịa): {d}")
    for n in set(re.findall(r"\d{3,}", text)):
        if n not in nums_ok:
            errs.append(f"số không có trong đầu vào (nghi bịa): {n}")
    for v in nums:
        val = v * 100 if 0 < v <= 1 else v  # ô phần trăm lưu dạng 0–1
        if float(val).is_integer() and len(str(int(val))) >= 3:
            sv = str(int(val))
            if sv not in nums_ok and sv not in expect.get("allowed_derived", []):
                errs.append(f"số trong ô không có trong đầu vào: {sv}")
    for label in expect.get("blank_labels", []):
        lines = [ln for ln in text.split("\n") if ln.strip().startswith(label)]
        if not lines:
            errs.append(f"không thấy nhãn cần chừa chỗ: {label!r}")
        elif not any(re.search(r"\.{4,}|…", ln) or ln.strip() == label for ln in lines) and \
                not any(re.search(r"\.{4,}", t) for t in _following(text, label)):
            errs.append(f"nhãn {label!r} chưa chừa chỗ điền (dòng dấu chấm)")
    findings = results / f"{name}.findings.md"
    if expect.get("traps"):
        if not findings.is_file():
            errs.append("thiếu file ghi nhận phát hiện (findings.md)")
        else:
            ft = findings.read_text(encoding="utf-8")
            for t in expect["traps"]:
                if t["id"] not in ft:
                    errs.append(f"findings chưa ghi nhận: {t['id']}")
    return errs


def _following(text, label):
    lines = text.split("\n")
    for i, ln in enumerate(lines):
        if ln.strip().startswith(label):
            return lines[i + 1:i + 2]
    return []


def main():
    ap = argparse.ArgumentParser()
    root = Path(__file__).resolve().parents[1]
    ap.add_argument("--samples", default=root / "tests/samples")
    ap.add_argument("--outputs", default=root / "tests/outputs")
    ap.add_argument("--results", default=root / "tests/results")
    a = ap.parse_args()
    samples, outputs, results = Path(a.samples), Path(a.outputs), Path(a.results)
    failed = 0
    summary = {}
    for ep in sorted(samples.glob("*.expect.json")):
        expect = json.loads(ep.read_text(encoding="utf-8"))
        log = []
        errs = check_one(expect, samples, outputs, results, log)
        summary[expect["skill"]] = {"dat": not errs, "loi": errs}
        if log:
            summary[expect["skill"]]["ghi_chu"] = [x.strip() for x in log]
        failed += bool(errs)
    print(json.dumps(summary, ensure_ascii=False, indent=1))
    return 1 if failed else 0


if __name__ == "__main__":
    sys.exit(main())
